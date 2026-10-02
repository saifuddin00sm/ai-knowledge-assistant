"""Document parsers. Each returns pages so PDF page numbers survive into metadata."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

from app.core.errors import InvalidRequestError, UnsupportedMediaTypeError

SUPPORTED_EXTENSIONS = ("pdf", "docx", "txt", "md")

_WHITESPACE_RUN = re.compile(r"[ \t]{2,}")
_BLANK_LINES = re.compile(r"\n{3,}")

# Non-breaking space -> space; zero-width space and BOM -> removed. Referenced by
# code point so the characters themselves never appear in this source file.
_INVISIBLES = str.maketrans({0x00A0: " ", 0x200B: None, 0xFEFF: None})


@dataclass(slots=True)
class ParsedPage:
    """One unit of source text. `page` is 1-based for PDFs, None otherwise."""

    text: str
    page: int | None = None


@dataclass(slots=True)
class ParsedDocument:
    pages: list[ParsedPage] = field(default_factory=list)
    page_count: int | None = None

    @property
    def text(self) -> str:
        return "\n\n".join(page.text for page in self.pages)

    @property
    def is_empty(self) -> bool:
        return not any(page.text.strip() for page in self.pages)


def normalize_text(text: str) -> str:
    """Collapse the noise PDF extraction leaves behind, keep paragraph breaks."""
    text = text.replace("\r\n", "\n").replace("\r", "\n").translate(_INVISIBLES)
    text = _WHITESPACE_RUN.sub(" ", text)
    text = "\n".join(line.rstrip() for line in text.split("\n"))
    return _BLANK_LINES.sub("\n\n", text).strip()


def extension_of(filename: str) -> str:
    _, _, suffix = filename.rpartition(".")
    return suffix.lower() if suffix and suffix != filename else ""


def parse_pdf(data: bytes) -> ParsedDocument:
    import fitz  # PyMuPDF

    try:
        document = fitz.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise InvalidRequestError(f"Could not open PDF: {exc}") from exc

    pages: list[ParsedPage] = []
    try:
        for number, page in enumerate(document, start=1):
            text = normalize_text(page.get_text("text"))
            if text:
                pages.append(ParsedPage(text=text, page=number))
        total = document.page_count
    finally:
        document.close()
    return ParsedDocument(pages=pages, page_count=total)


def parse_docx(data: bytes) -> ParsedDocument:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise InvalidRequestError(f"Could not open DOCX: {exc}") from exc

    blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))
    text = normalize_text("\n\n".join(blocks))
    return ParsedDocument(pages=[ParsedPage(text=text)] if text else [], page_count=None)


def parse_plain_text(data: bytes) -> ParsedDocument:
    text = normalize_text(data.decode("utf-8", errors="replace"))
    return ParsedDocument(pages=[ParsedPage(text=text)] if text else [], page_count=None)


def parse(filename: str, data: bytes) -> ParsedDocument:
    """Dispatch on file extension. Raises for anything unsupported or unreadable."""
    extension = extension_of(filename)
    if extension == "pdf":
        return parse_pdf(data)
    if extension == "docx":
        return parse_docx(data)
    if extension in {"txt", "md"}:
        return parse_plain_text(data)
    raise UnsupportedMediaTypeError(
        f"Unsupported file type '.{extension or filename}'. "
        f"Allowed: {', '.join(SUPPORTED_EXTENSIONS)}."
    )
